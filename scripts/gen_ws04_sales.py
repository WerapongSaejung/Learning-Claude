# -*- coding: utf-8 -*-
"""
สร้างไฟล์ข้อมูลตัวอย่างสำหรับ Workshop #4 — "ร้านขายอุปกรณ์คอมพิวเตอร์"
ยอดขาย 1,000 รายการ · สินค้า 100 ชนิด (100 barcode) · ตลอดปี 2568

ไฟล์นี้เป็นตัวสำรองสำหรับวิทยากร: ถ้าหน้างาน Cowork สุ่มข้อมูลไม่สำเร็จ
หรือใช้เวลานานเกินไป ให้วางไฟล์นี้ลงโฟลเดอร์ workshop แล้วข้ามไป Prompt 2 ได้เลย

ข้อมูลทั้งหมดเป็นข้อมูลสมมติ ไม่ใช่ข้อมูลของบริษัทใด
รันซ้ำได้ผลเดิมทุกครั้ง (seed คงที่)
"""
import random
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SEED = 20260824
ROWS = 1000
OUT = "workshop/WS04-sales-sample.xlsx"
FONT = "Arial"

random.seed(SEED)

# ---------------------------------------------------------------- สินค้า 100 ชนิด
# (หมวดหมู่, ยี่ห้อ, ชื่อรุ่น, ราคาขายปลีก, น้ำหนักความถี่ในการขาย)
CATALOG = [
    # การ์ดจอ 8
    ("การ์ดจอ", "NVIDIA", "GeForce RTX 4060 8GB", 12900, 6),
    ("การ์ดจอ", "NVIDIA", "GeForce RTX 4060 Ti 16GB", 19900, 4),
    ("การ์ดจอ", "NVIDIA", "GeForce RTX 4070 Super 12GB", 27900, 3),
    ("การ์ดจอ", "NVIDIA", "GeForce RTX 4080 Super 16GB", 45900, 2),
    ("การ์ดจอ", "AMD", "Radeon RX 7600 8GB", 10500, 4),
    ("การ์ดจอ", "AMD", "Radeon RX 7700 XT 12GB", 17900, 3),
    ("การ์ดจอ", "AMD", "Radeon RX 7800 XT 16GB", 22900, 2),
    ("การ์ดจอ", "Intel", "Arc A750 8GB", 8900, 2),
    # ซีพียู 8
    ("ซีพียู", "Intel", "Core i3-14100F", 3990, 6),
    ("ซีพียู", "Intel", "Core i5-14400F", 7290, 8),
    ("ซีพียู", "Intel", "Core i5-14600K", 11900, 5),
    ("ซีพียู", "Intel", "Core i7-14700K", 16900, 3),
    ("ซีพียู", "Intel", "Core i9-14900K", 23900, 2),
    ("ซีพียู", "AMD", "Ryzen 5 7600", 7590, 6),
    ("ซีพียู", "AMD", "Ryzen 7 7800X3D", 15900, 4),
    ("ซีพียู", "AMD", "Ryzen 9 7950X", 21900, 2),
    # เมนบอร์ด 8
    ("เมนบอร์ด", "ASUS", "PRIME B760M-A", 4290, 6),
    ("เมนบอร์ด", "ASUS", "ROG STRIX B650-A", 8900, 4),
    ("เมนบอร์ด", "ASUS", "ROG MAXIMUS Z790 HERO", 13500, 1),
    ("เมนบอร์ด", "MSI", "PRO B760M-P", 3590, 7),
    ("เมนบอร์ด", "MSI", "MAG B650 TOMAHAWK", 7490, 4),
    ("เมนบอร์ด", "Gigabyte", "B760M DS3H", 3390, 6),
    ("เมนบอร์ด", "Gigabyte", "X670 AORUS ELITE", 9900, 2),
    ("เมนบอร์ด", "ASRock", "A620M-HDV", 2490, 5),
    # แรม 8
    ("แรม", "Kingston", "FURY Beast DDR4 8GB 3200", 1190, 9),
    ("แรม", "Kingston", "FURY Beast DDR4 16GB 3200", 1890, 10),
    ("แรม", "Kingston", "FURY Beast DDR5 16GB 5600", 2490, 8),
    ("แรม", "Kingston", "FURY Renegade DDR5 32GB 6000", 4990, 4),
    ("แรม", "Corsair", "Vengeance DDR5 16GB 5600", 2690, 7),
    ("แรม", "Corsair", "Vengeance RGB DDR5 32GB 6000", 5490, 4),
    ("แรม", "G.Skill", "Trident Z5 DDR5 32GB 6400", 6900, 2),
    ("แรม", "G.Skill", "Ripjaws V DDR4 16GB 3600", 1990, 6),
    # SSD / HDD 10
    ("SSD / HDD", "Samsung", "970 EVO Plus NVMe 500GB", 1890, 8),
    ("SSD / HDD", "Samsung", "980 PRO NVMe 1TB", 3690, 7),
    ("SSD / HDD", "Samsung", "990 PRO NVMe 2TB", 7490, 3),
    ("SSD / HDD", "WD", "Blue SN580 NVMe 500GB", 1490, 9),
    ("SSD / HDD", "WD", "Black SN770 NVMe 1TB", 2990, 6),
    ("SSD / HDD", "WD", "Blue HDD 2TB 7200RPM", 1990, 5),
    ("SSD / HDD", "Seagate", "Barracuda HDD 4TB", 3290, 4),
    ("SSD / HDD", "Seagate", "IronWolf NAS HDD 8TB", 8500, 2),
    ("SSD / HDD", "Crucial", "P3 Plus NVMe 1TB", 2290, 7),
    ("SSD / HDD", "Kingston", "NV2 NVMe 500GB", 1290, 8),
    # จอมอนิเตอร์ 10
    ("จอมอนิเตอร์", "LG", "24MK600M 24\" IPS 75Hz", 3290, 8),
    ("จอมอนิเตอร์", "LG", "27GP850 27\" QHD 180Hz", 11900, 4),
    ("จอมอนิเตอร์", "LG", "34WP65C 34\" UltraWide", 13900, 2),
    ("จอมอนิเตอร์", "Samsung", "S24C310 24\" IPS 75Hz", 2990, 9),
    ("จอมอนิเตอร์", "Samsung", "Odyssey G5 27\" QHD 165Hz", 8900, 5),
    ("จอมอนิเตอร์", "Samsung", "Odyssey G7 32\" 4K 165Hz", 18900, 2),
    ("จอมอนิเตอร์", "AOC", "24G2SPE 24\" 165Hz", 4590, 7),
    ("จอมอนิเตอร์", "AOC", "27G2SPU 27\" 165Hz", 6290, 5),
    ("จอมอนิเตอร์", "BenQ", "GW2480 24\" IPS", 3690, 6),
    ("จอมอนิเตอร์", "ViewSonic", "VA2432 24\" IPS", 3190, 6),
    # คีย์บอร์ด 8
    ("คีย์บอร์ด", "Logitech", "K120 USB", 390, 12),
    ("คีย์บอร์ด", "Logitech", "MX Keys S Wireless", 4290, 4),
    ("คีย์บอร์ด", "Logitech", "G413 TKL SE Mechanical", 2190, 6),
    ("คีย์บอร์ด", "Razer", "BlackWidow V4 X", 4990, 3),
    ("คีย์บอร์ด", "Keychron", "K2 Pro Wireless Mechanical", 5900, 3),
    ("คีย์บอร์ด", "SteelSeries", "Apex 3 TKL", 2490, 4),
    ("คีย์บอร์ด", "Corsair", "K55 RGB PRO", 1990, 5),
    ("คีย์บอร์ด", "Nubwo", "NK-30 Gaming", 590, 9),
    # เมาส์ 8
    ("เมาส์", "Logitech", "B100 USB", 250, 14),
    ("เมาส์", "Logitech", "M331 Silent Wireless", 690, 11),
    ("เมาส์", "Logitech", "MX Master 3S", 3890, 5),
    ("เมาส์", "Logitech", "G304 Lightspeed", 1290, 8),
    ("เมาส์", "Razer", "DeathAdder V3", 2790, 4),
    ("เมาส์", "SteelSeries", "Rival 3", 1190, 5),
    ("เมาส์", "Corsair", "Katar Pro XT", 990, 6),
    ("เมาส์", "Nubwo", "NM-81 Gaming", 350, 10),
    # หูฟัง / ลำโพง 8
    ("หูฟัง / ลำโพง", "Logitech", "H390 USB Headset", 990, 8),
    ("หูฟัง / ลำโพง", "Logitech", "Z313 Speaker 2.1", 1290, 6),
    ("หูฟัง / ลำโพง", "HyperX", "Cloud Stinger 2", 1890, 6),
    ("หูฟัง / ลำโพง", "HyperX", "Cloud III Wireless", 5490, 3),
    ("หูฟัง / ลำโพง", "SteelSeries", "Arctis Nova 1", 2990, 4),
    ("หูฟัง / ลำโพง", "Razer", "BlackShark V2 X", 2290, 4),
    ("หูฟัง / ลำโพง", "Edifier", "R1280DB Bookshelf", 5900, 2),
    ("หูฟัง / ลำโพง", "Nubwo", "N2 Gaming Headset", 590, 9),
    # เคส / พาวเวอร์ซัพพลาย 8
    ("เคส / เพาเวอร์", "NZXT", "H5 Flow ATX Case", 3290, 4),
    ("เคส / เพาเวอร์", "Lian Li", "Lancool 216 Case", 3990, 3),
    ("เคส / เพาเวอร์", "Cooler Master", "MasterBox Q300L", 1690, 6),
    ("เคส / เพาเวอร์", "Corsair", "4000D Airflow", 3590, 4),
    ("เคส / เพาเวอร์", "Corsair", "RM750e 750W 80+ Gold", 3490, 5),
    ("เคส / เพาเวอร์", "Seasonic", "FOCUS GX-850 850W", 4890, 3),
    ("เคส / เพาเวอร์", "Cooler Master", "MWE Bronze 650W", 1990, 6),
    ("เคส / เพาเวอร์", "Antec", "CSK 550W 80+ Bronze", 1590, 5),
    # อุปกรณ์เครือข่าย 6
    ("อุปกรณ์เครือข่าย", "TP-Link", "Archer C6 Router AC1200", 990, 8),
    ("อุปกรณ์เครือข่าย", "TP-Link", "Archer AX53 Wi-Fi 6", 1890, 6),
    ("อุปกรณ์เครือข่าย", "TP-Link", "TL-SG1008D Switch 8 Port", 690, 7),
    ("อุปกรณ์เครือข่าย", "ASUS", "RT-AX57 Wi-Fi 6 Router", 2790, 4),
    ("อุปกรณ์เครือข่าย", "Ubiquiti", "UniFi U6 Lite AP", 4900, 2),
    ("อุปกรณ์เครือข่าย", "D-Link", "DIR-825 Router", 1290, 4),
    # สายเคเบิล / อะแดปเตอร์ 10
    ("สาย / อะแดปเตอร์", "Ugreen", "HDMI 2.1 Cable 2m", 390, 16),
    ("สาย / อะแดปเตอร์", "Ugreen", "USB-C to HDMI Adapter", 590, 12),
    ("สาย / อะแดปเตอร์", "Ugreen", "USB-C Hub 6-in-1", 890, 10),
    ("สาย / อะแดปเตอร์", "Orico", "USB 3.0 Hub 4 Port", 350, 13),
    ("สาย / อะแดปเตอร์", "Orico", "M.2 NVMe Enclosure", 690, 7),
    ("สาย / อะแดปเตอร์", "Anker", "PowerLine USB-C 1.8m", 490, 11),
    ("สาย / อะแดปเตอร์", "Anker", "65W GaN Charger", 890, 8),
    ("สาย / อะแดปเตอร์", "Vention", "DisplayPort 1.4 Cable 2m", 290, 12),
    ("สาย / อะแดปเตอร์", "Vention", "CAT6 LAN Cable 5m", 190, 14),
    ("สาย / อะแดปเตอร์", "Generic", "Cable Tie / Velcro Pack", 90, 15),
]
assert len(CATALOG) == 100, len(CATALOG)


def ean13(seq):
    """สร้างบาร์โค้ด EAN-13 ขึ้นต้น 885 (รหัสประเทศไทย) พร้อม check digit ที่ถูกต้อง"""
    body = "885" + str(1000000 + seq).zfill(9)[:9]
    total = sum(int(d) * (3 if i % 2 else 1) for i, d in enumerate(body))
    return body + str((10 - total % 10) % 10)


PRODUCTS = []
for i, (cat, brand, name, price, weight) in enumerate(CATALOG, start=1):
    PRODUCTS.append({
        "barcode": ean13(i),
        "cat": cat,
        "brand": brand,
        "name": "%s %s" % (brand, name),
        "price": price,
        "weight": weight,
    })

# ---------------------------------------------------------------- ช่องทาง / พนักงาน
CHANNELS = ["หน้าร้าน", "เว็บไซต์ร้าน", "Shopee", "Lazada", "ขายส่ง (B2B)"]
STAFF = ["ณัฐพงษ์", "ปิยะดา", "ชัยวัฒน์", "อรทัย", "สุเมธ", "กนกวรรณ"]

START = date(2025, 1, 1)
END = date(2025, 12, 31)
DAYS = (END - START).days + 1


def month_factor(m):
    """ฤดูกาลขาย: ต้นปีเงียบ · กลางปีทรงตัว · พ.ย.-ธ.ค. พุ่ง (11.11 / 12.12 / ปีใหม่)"""
    return {1: .78, 2: .72, 3: .85, 4: .80, 5: .92, 6: .95,
            7: 1.00, 8: 1.02, 9: 1.05, 10: 1.15, 11: 1.55, 12: 1.45}[m]


def channel_weights(m):
    """ออนไลน์โตขึ้นเรื่อยๆ ตลอดปี หน้าร้านค่อยๆ ลด"""
    online_shift = (m - 1) * 0.9
    return [
        max(6.0, 34 - online_shift * 1.4),   # หน้าร้าน
        14 + online_shift * 0.35,            # เว็บไซต์ร้าน
        22 + online_shift * 0.65,            # Shopee
        18 + online_shift * 0.45,            # Lazada
        12,                                  # ขายส่ง (B2B)
    ]


def pick_day():
    """สุ่มวัน โดยถ่วงน้ำหนักตามฤดูกาล และเสาร์-อาทิตย์ขายดีกว่าวันธรรมดา"""
    while True:
        d = START + timedelta(days=random.randrange(DAYS))
        w = month_factor(d.month) * (1.25 if d.weekday() >= 5 else 1.0)
        if random.random() < w / 2.0:
            return d


def qty_for(price):
    if price >= 10000:
        return 1
    if price >= 3000:
        return random.choices([1, 2], weights=[88, 12])[0]
    if price >= 1000:
        return random.choices([1, 2, 3], weights=[74, 20, 6])[0]
    return random.choices([1, 2, 3, 4, 5], weights=[46, 26, 15, 8, 5])[0]


rows = []
bill_no = 68000
for _ in range(ROWS):
    d = pick_day()

    # ── สินค้า: ของถูกขายบ่อยกว่า และมีสินค้า 1 ตัวที่เลิกขายกลางปี ──
    while True:
        p = random.choices(PRODUCTS, weights=[x["weight"] for x in PRODUCTS])[0]
        # Intel Arc A750 เลิกทำตลาดตั้งแต่ ก.ค. เป็นต้นไป
        if p["name"].endswith("Arc A750 8GB") and d.month >= 7:
            continue
        break

    ch = random.choices(CHANNELS, weights=channel_weights(d.month))[0]
    qty = qty_for(p["price"])
    if ch == "ขายส่ง (B2B)":
        qty = max(qty, random.choices([2, 3, 5, 8, 10], weights=[30, 26, 22, 12, 10])[0])

    staff = random.choice(STAFF) if ch in ("หน้าร้าน", "ขายส่ง (B2B)") else random.choice(STAFF[:4])

    unit = p["price"]
    # ราคาแพลตฟอร์มมาร์เก็ตเพลสตั้งสูงกว่านิดหน่อยเพื่อชดเชยค่า GP
    if ch in ("Shopee", "Lazada"):
        unit = round(unit * random.uniform(1.01, 1.05) / 10) * 10

    subtotal = unit * qty
    # ── ส่วนลด: ปกติไม่ค่อยลด · ขายส่งลดเสมอ · "ชัยวัฒน์" ลดหนักผิดปกติ ──
    if ch == "ขายส่ง (B2B)":
        rate = random.uniform(.05, .12)
    elif staff == "ชัยวัฒน์":
        rate = random.choices([0, random.uniform(.05, .15)], weights=[35, 65])[0]
    else:
        rate = random.choices([0, random.uniform(.02, .07)], weights=[76, 24])[0]
    if d.month in (11, 12) and rate == 0:
        rate = random.choices([0, random.uniform(.03, .10)], weights=[45, 55])[0]
    discount = round(subtotal * rate / 10) * 10

    bill_no += random.randrange(1, 4)
    rows.append({
        "date": d,
        "bill": "INV-68-%05d" % bill_no,
        "barcode": p["barcode"],
        "name": p["name"],
        "cat": p["cat"],
        "brand": p["brand"],
        "qty": qty,
        "unit": unit,
        "discount": discount,
        "channel": ch,
        "staff": staff,
    })

rows.sort(key=lambda r: (r["date"], r["bill"]))

# ---------------------------------------------------------------- เขียนไฟล์
HEADERS = ["วันที่", "เลขที่ใบเสร็จ", "บาร์โค้ด", "ชื่อสินค้า", "หมวดหมู่", "ยี่ห้อ",
           "จำนวน", "ราคาต่อหน่วย", "ส่วนลด", "ยอดสุทธิ", "ช่องทางขาย", "พนักงานขาย"]
WIDTHS = [12, 15, 16, 34, 20, 13, 9, 15, 11, 14, 15, 14]

wb = Workbook()
ws = wb.active
ws.title = "ยอดขาย"

head_fill = PatternFill("solid", fgColor="1F3864")
head_font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
body_font = Font(name=FONT, size=11)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c, h in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill, cell.font, cell.border = head_fill, head_font, border
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 24

for i, r in enumerate(rows, start=2):
    vals = [r["date"], r["bill"], r["barcode"], r["name"], r["cat"], r["brand"],
            r["qty"], r["unit"], r["discount"], "=G%d*H%d-I%d" % (i, i, i),
            r["channel"], r["staff"]]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font, cell.border = body_font, border
    ws.cell(row=i, column=1).number_format = "dd/mm/yyyy"
    ws.cell(row=i, column=3).alignment = Alignment(horizontal="left")
    ws.cell(row=i, column=7).number_format = "#,##0"
    for c in (8, 9, 10):
        ws.cell(row=i, column=c).number_format = "#,##0.00"

for c, w in enumerate(WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:L%d" % (len(rows) + 1)

# ---- ชีตอธิบายคอลัมน์ (ไม่มีเฉลย — เฉลยอยู่ใน workshop/README.md สำหรับวิทยากร) ----
info = wb.create_sheet("อ่านก่อน")
info.column_dimensions["A"].width = 20
info.column_dimensions["B"].width = 82
lines = [
    ("ข้อมูลชุดนี้คืออะไร", "ยอดขายสมมติของร้านขายอุปกรณ์คอมพิวเตอร์ ใช้สำหรับฝึกวิเคราะห์ข้อมูลใน Workshop #4"),
    ("", "เป็นข้อมูลที่สร้างขึ้นเองทั้งหมด ไม่ใช่ข้อมูลของบริษัทหรือร้านค้าจริง"),
    ("ช่วงเวลา", "1 มกราคม – 31 ธันวาคม 2568"),
    ("จำนวนรายการ", "1,000 รายการ · สินค้า 100 ชนิด (100 บาร์โค้ด)"),
    ("", ""),
    ("คำอธิบายคอลัมน์", ""),
    ("วันที่", "วันที่ขาย"),
    ("เลขที่ใบเสร็จ", "เลขที่ใบเสร็จ 1 ใบ = 1 รายการในตารางนี้"),
    ("บาร์โค้ด", "รหัส EAN-13 ของสินค้า (ขึ้นต้น 885 = รหัสประเทศไทย)"),
    ("ชื่อสินค้า", "ยี่ห้อ + ชื่อรุ่น"),
    ("หมวดหมู่", "12 หมวด เช่น การ์ดจอ ซีพียู แรม จอมอนิเตอร์ สาย/อะแดปเตอร์"),
    ("ยี่ห้อ", "ยี่ห้อสินค้า"),
    ("จำนวน", "จำนวนชิ้นที่ขายในรายการนั้น"),
    ("ราคาต่อหน่วย", "ราคาขายต่อชิ้น (บาท) — ราคาบนมาร์เก็ตเพลสตั้งสูงกว่าหน้าร้านเล็กน้อย"),
    ("ส่วนลด", "ส่วนลดเป็นจำนวนเงิน (บาท) ของทั้งรายการ"),
    ("ยอดสุทธิ", "สูตร = จำนวน × ราคาต่อหน่วย − ส่วนลด"),
    ("ช่องทางขาย", "หน้าร้าน · เว็บไซต์ร้าน · Shopee · Lazada · ขายส่ง (B2B)"),
    ("พนักงานขาย", "ชื่อพนักงานที่ปิดการขาย"),
]
lab = Font(name=FONT, size=11, bold=True)
val = Font(name=FONT, size=11)
for i, (a, b) in enumerate(lines, start=1):
    ca = info.cell(row=i, column=1, value=a)
    cb = info.cell(row=i, column=2, value=b)
    ca.font, cb.font = lab, val
    ca.alignment = Alignment(vertical="top")
    cb.alignment = Alignment(vertical="top", wrap_text=True)

wb.save(OUT)
print("wrote %s  rows=%d  products=%d" % (OUT, len(rows), len(PRODUCTS)))
