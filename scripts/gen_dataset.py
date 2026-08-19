# -*- coding: utf-8 -*-
"""สร้างไฟล์ Excel ข้อมูลสมมติสำหรับ workshop "ใช้ AI ทำงานจริงด้วย Claude"

ข้อมูลทั้งหมดแต่งขึ้นเพื่อการฝึกอบรม ชื่อบริษัทและชื่อผลิตภัณฑ์ใช้ของจริง
เพื่อให้ผู้เรียนรู้สึกใกล้ตัว แต่ตัวเลข ลูกค้า และเหตุการณ์ไม่มีอยู่จริง

รันแล้วได้ 3 ไฟล์ใน dataset/
  TP-01-sales-monthly.xlsx   ยอดขายรายเดือน 720 แถว (ซ่อนความผิดปกติ 3 จุด)
  TP-02-product-cost.xlsx    ต้นทุนและกำไรขั้นต้นรายผลิตภัณฑ์
  TP-03-complaint-log.xlsx   ทะเบียนข้อร้องเรียนผลิตภัณฑ์ 80 รายการ
"""
import os
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'dataset')
FONT = 'Tahoma'          # อ่านภาษาไทยได้ทุกเครื่อง และเป็นฟอนต์มาตรฐานเอกสารธุรกิจไทย
DISCLAIMER = 'ข้อมูลสมมติเพื่อการฝึกอบรม — ไม่ใช่ข้อมูลจริงของบริษัท ห้ามนำไปใช้อ้างอิงในงานจริง'

HEAD_FILL = PatternFill('solid', fgColor='1F3864')
NOTE_FILL = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ---------------------------------------------------------------- ข้อมูลตั้งต้น

MONTHS = ['ม.ค.', 'ก.พ.', 'มี.ค.', 'เม.ย.', 'พ.ค.', 'มิ.ย.',
          'ก.ค.', 'ส.ค.', 'ก.ย.', 'ต.ค.', 'พ.ย.', 'ธ.ค.']

# (ชื่อผลิตภัณฑ์, กลุ่ม, หน่วยนับ, ราคาอ้างอิง/หน่วย, จำนวนหน่วยฐานต่อเดือน)
PRODUCTS = [
    ('Ceftriaxone T P',              'ยาฉีดปราศจากเชื้อ — Cephalosporin', 'vial',          42.0, 62000),
    ('Cefazillin',                   'ยาฉีดปราศจากเชื้อ — Cephalosporin', 'vial',          38.0, 41000),
    ('Clatacef',                     'ยาฉีดปราศจากเชื้อ — Cephalosporin', 'vial',          55.0, 23000),
    ('Sterile Ampicillin Sodium T P', 'ยาฉีดปราศจากเชื้อ',                 'vial',          26.0, 48000),
    ('Gentamicin T P',               'ยาฉีดปราศจากเชื้อ',                 'ampoule',       14.0, 55000),
    ('Kanamycin Sulfate Injection T P', 'ยาฉีดปราศจากเชื้อ',              'vial',          31.0, 19000),
    ('Furosemide Injection T P',     'ยาฉีดปราศจากเชื้อ',                 'ampoule',       12.0, 67000),
    ('Vitamin C Injection T P',      'ยาฉีดปราศจากเชื้อ',                 'ampoule',        9.0, 88000),
    ('B-100 Complex Injection T P',  'ยาฉีดปราศจากเชื้อ',                 'ampoule',       16.0, 52000),
    ('Paramol T P',                  'ยาเม็ด',                            'กล่อง 100 เม็ด', 48.0, 34000),
    ('Tramadol T P',                 'ยาเม็ด',                            'กล่อง 100 เม็ด', 120.0, 12000),
    ('K-CIL',                        'ยาเม็ด',                            'กล่อง 100 เม็ด', 95.0, 15000),
    ('Dexton',                       'ยาเม็ด',                            'กล่อง 100 เม็ด', 68.0, 21000),
    ('Medeton',                      'ยาใช้ภายนอก',                       'ขวด',           35.0, 26000),
    ('Transic',                      'ยาใช้ภายนอก',                       'ขวด',           44.0, 18000),
]

# ช่องทางจำหน่าย: (ชื่อ, สัดส่วนหน่วยขาย, ตัวคูณราคาเทียบราคาอ้างอิง)
CHANNELS = [
    ('โรงพยาบาลรัฐ',    0.42, 0.90),   # ได้จากการประมูล ราคาต่ำกว่าราคาอ้างอิง
    ('โรงพยาบาลเอกชน',  0.18, 1.05),
    ('ร้านขายยา',       0.22, 1.00),
    ('ส่งออก',          0.18, 0.78),   # ราคาต่อหน่วยต่ำสุด แต่ปริมาณต่อคำสั่งซื้อสูง
]

SEASONALITY = [0.92, 0.95, 1.05, 0.98, 1.02, 1.06, 1.00, 0.97, 1.03, 1.08, 1.06, 0.88]


def anomaly_factor(product, channel, month):
    """ความผิดปกติ 3 จุดที่ตั้งใจซ่อนไว้ให้ผู้เรียนหาเจอใน M2 (เฉลยใน dataset/README.md)"""
    f = 1.0
    # จุดที่ 1 — เสียงานประมูลโรงพยาบาลรัฐของ Ceftriaxone ตั้งแต่ ก.ค.
    if product == 'Ceftriaxone T P' and channel == 'โรงพยาบาลรัฐ' and month >= 7:
        f *= 0.18
    # จุดที่ 2 — ยอดส่งออกโตต่อเนื่องตั้งแต่ เม.ย. ไปกลบยอดที่หายไปในภาพรวม
    if channel == 'ส่งออก' and month >= 4:
        f *= 1.0 + 0.07 * (month - 3)
    # จุดที่ 3 — Gentamicin ถูกระงับจำหน่ายจากข้อร้องเรียนเรื่องอนุภาคแขวนลอย
    if product == 'Gentamicin T P':
        if month == 8:
            f *= 0.45
        elif month == 9:
            f *= 0.0
    return f


def style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BOX
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_note_sheet(wb, lines):
    ws = wb.create_sheet('หมายเหตุ')
    ws['A1'] = 'หมายเหตุประกอบไฟล์'
    ws['A1'].font = Font(name=FONT, size=13, bold=True)
    r = 3
    for line in [DISCLAIMER] + lines:
        c = ws.cell(row=r, column=1, value=line)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if r == 3:
            c.fill = NOTE_FILL
            c.font = Font(name=FONT, size=10, bold=True)
        ws.row_dimensions[r].height = 30
        r += 2
    ws.column_dimensions['A'].width = 110
    return ws


# ------------------------------------------------------- TP-01 ยอดขายรายเดือน

def build_sales():
    rng = random.Random(25680101)
    wb = Workbook()
    ws = wb.active
    ws.title = 'ยอดขายรายเดือน'

    headers = ['เดือน', 'ลำดับเดือน', 'ผลิตภัณฑ์', 'กลุ่มผลิตภัณฑ์', 'ช่องทางจำหน่าย',
               'หน่วยนับ', 'จำนวนที่ขาย', 'ราคาต่อหน่วย (บาท)', 'ยอดขาย (บาท)']
    ws.append(headers)

    rows = 0
    for m in range(1, 13):
        for name, group, unit, price, base in PRODUCTS:
            for ch, share, price_mult in CHANNELS:
                units = base * share * SEASONALITY[m - 1] * anomaly_factor(name, ch, m)
                units *= 1.0 + rng.uniform(-0.06, 0.06)
                units = int(round(units))
                unit_price = round(price * price_mult, 2)
                ws.append([MONTHS[m - 1], m, name, group, ch, unit, units, unit_price, None])
                rows += 1
                r = rows + 1
                ws.cell(row=r, column=9, value='=G%d*H%d' % (r, r))

    style_header(ws)
    for r in range(2, rows + 2):
        for c in range(1, 10):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BOX
        ws.cell(row=r, column=7).number_format = '#,##0'
        ws.cell(row=r, column=8).number_format = '#,##0.00'
        ws.cell(row=r, column=9).number_format = '#,##0.00'
    autosize(ws, [10, 12, 30, 34, 18, 16, 14, 18, 18])

    # สรุปรายเดือน — ใช้สูตรทั้งหมด เพื่อให้แก้ข้อมูลต้นทางแล้วตัวเลขขยับตาม
    s = wb.create_sheet('สรุปรายเดือน')
    s.append(['เดือน', 'ลำดับเดือน', 'ไตรมาส', 'ยอดขายรวม (บาท)',
              'ยอดขาย รพ.รัฐ', 'ยอดขาย รพ.เอกชน', 'ยอดขาย ร้านขายยา', 'ยอดขาย ส่งออก'])
    last = rows + 1
    for m in range(1, 13):
        r = m + 1
        s.append([MONTHS[m - 1], m, 'Q%d' % ((m - 1) // 3 + 1),
                  '=SUMIFS(ยอดขายรายเดือน!$I$2:$I$%d,ยอดขายรายเดือน!$B$2:$B$%d,$B%d)' % (last, last, r),
                  '=SUMIFS(ยอดขายรายเดือน!$I$2:$I$%d,ยอดขายรายเดือน!$B$2:$B$%d,$B%d,ยอดขายรายเดือน!$E$2:$E$%d,"โรงพยาบาลรัฐ")' % (last, last, r, last),
                  '=SUMIFS(ยอดขายรายเดือน!$I$2:$I$%d,ยอดขายรายเดือน!$B$2:$B$%d,$B%d,ยอดขายรายเดือน!$E$2:$E$%d,"โรงพยาบาลเอกชน")' % (last, last, r, last),
                  '=SUMIFS(ยอดขายรายเดือน!$I$2:$I$%d,ยอดขายรายเดือน!$B$2:$B$%d,$B%d,ยอดขายรายเดือน!$E$2:$E$%d,"ร้านขายยา")' % (last, last, r, last),
                  '=SUMIFS(ยอดขายรายเดือน!$I$2:$I$%d,ยอดขายรายเดือน!$B$2:$B$%d,$B%d,ยอดขายรายเดือน!$E$2:$E$%d,"ส่งออก")' % (last, last, r, last)])
    s.append(['รวมทั้งปี', None, None,
              '=SUM(D2:D13)', '=SUM(E2:E13)', '=SUM(F2:F13)', '=SUM(G2:G13)', '=SUM(H2:H13)'])
    style_header(s)
    for r in range(2, 15):
        for c in range(1, 9):
            cell = s.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, bold=(r == 14))
            cell.border = BOX
            if c >= 4:
                cell.number_format = '#,##0'
    autosize(s, [12, 12, 10, 20, 18, 20, 20, 18])

    add_note_sheet(wb, [
        'ชีต "ยอดขายรายเดือน" คือข้อมูลระดับรายการ 720 แถว (12 เดือน x 15 ผลิตภัณฑ์ x 4 ช่องทาง) ปีบัญชี 2568',
        'ชีต "สรุปรายเดือน" คำนวณด้วยสูตร SUMIFS จากชีตแรก แก้ข้อมูลต้นทางแล้วตัวเลขสรุปขยับตามเอง',
        'ราคาต่อหน่วยต่างกันตามช่องทาง: โรงพยาบาลรัฐ 90% ของราคาอ้างอิง (ได้จากการประมูล), โรงพยาบาลเอกชน 105%, ร้านขายยา 100%, ส่งออก 78%',
        'ไฟล์นี้ใช้ประกอบโมดูล M2 (ทำงานกับไฟล์) และ M3 (สร้างไฟล์ออกมาจริง)',
    ])
    wb.save(os.path.join(OUT, 'TP-01-sales-monthly.xlsx'))
    return rows


# --------------------------------------------------- TP-02 ต้นทุนรายผลิตภัณฑ์

def build_cost():
    rng = random.Random(25680202)
    wb = Workbook()
    ws = wb.active
    ws.title = 'ต้นทุนรายผลิตภัณฑ์'
    ws.append(['ผลิตภัณฑ์', 'กลุ่มผลิตภัณฑ์', 'หน่วยนับ',
               'ต้นทุนวัตถุดิบ (บาท/หน่วย)', 'ต้นทุนบรรจุภัณฑ์ (บาท/หน่วย)',
               'ค่าแรงและโสหุ้ยผลิต (บาท/หน่วย)', 'ต้นทุนรวม (บาท/หน่วย)',
               'ราคาอ้างอิง (บาท/หน่วย)', 'กำไรขั้นต้น (บาท/หน่วย)', 'อัตรากำไรขั้นต้น'])

    for i, (name, group, unit, price, _base) in enumerate(PRODUCTS, start=2):
        total_ratio = rng.uniform(0.52, 0.71)
        total_cost = price * total_ratio
        material = round(total_cost * rng.uniform(0.55, 0.66), 2)
        packing = round(total_cost * rng.uniform(0.10, 0.16), 2)
        overhead = round(total_cost - material - packing, 2)
        ws.append([name, group, unit, material, packing, overhead,
                   '=D%d+E%d+F%d' % (i, i, i), price,
                   '=H%d-G%d' % (i, i), '=IFERROR(I%d/H%d,"")' % (i, i)])

    style_header(ws)
    for r in range(2, len(PRODUCTS) + 2):
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BOX
            if 4 <= c <= 9:
                cell.number_format = '#,##0.00'
            if c == 10:
                cell.number_format = '0.0%'
    autosize(ws, [30, 34, 16, 22, 24, 26, 20, 20, 22, 18])

    ch = wb.create_sheet('ตัวคูณราคาตามช่องทาง')
    ch.append(['ช่องทางจำหน่าย', 'ตัวคูณราคาเทียบราคาอ้างอิง', 'หมายเหตุ'])
    notes = {
        'โรงพยาบาลรัฐ': 'ราคาจากการประมูล ต่ำกว่าราคาอ้างอิง',
        'โรงพยาบาลเอกชน': 'ราคาสูงสุด แต่ปริมาณต่อคำสั่งซื้อต่ำ',
        'ร้านขายยา': 'ใช้ราคาอ้างอิงเป็นหลัก',
        'ส่งออก': 'ราคาต่อหน่วยต่ำสุด แลกกับปริมาณต่อคำสั่งซื้อสูง',
    }
    for name, _share, mult in CHANNELS:
        ch.append([name, mult, notes[name]])
    style_header(ch)
    for r in range(2, 6):
        for c in range(1, 4):
            cell = ch.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BOX
        ch.cell(row=r, column=2).number_format = '0%'
    autosize(ch, [22, 30, 50])

    add_note_sheet(wb, [
        'ต้นทุนรวมต่อหน่วยและกำไรขั้นต้นคำนวณด้วยสูตรในไฟล์ ไม่ได้กรอกเป็นตัวเลขตายตัว',
        'อัตรากำไรขั้นต้นในชีตนี้คิดจากราคาอ้างอิง ยังไม่ได้คิดตัวคูณราคาตามช่องทาง',
        'ชีต "ตัวคูณราคาตามช่องทาง" คือกุญแจสำคัญของโจทย์ M3 — ยอดขายที่โตจากช่องทางส่งออกให้กำไรต่อหน่วยต่ำกว่าช่องทางอื่น',
        'ไฟล์นี้ใช้ประกอบโมดูล M3 (สร้างไฟล์ออกมาจริง) คู่กับ TP-01-sales-monthly.xlsx',
    ])
    wb.save(os.path.join(OUT, 'TP-02-product-cost.xlsx'))


# -------------------------------------------- TP-03 ทะเบียนข้อร้องเรียนสินค้า

COMPLAINT_TYPES = [
    ('คุณภาพทางกายภาพ', 'พบอนุภาคแขวนลอยในสารละลาย'),
    ('คุณภาพทางกายภาพ', 'ผงยาจับตัวเป็นก้อน'),
    ('คุณภาพทางกายภาพ', 'สีของสารละลายเปลี่ยนจากที่ระบุ'),
    ('บรรจุภัณฑ์', 'ฉลากพิมพ์เลื่อน อ่านเลขที่รุ่นผลิตไม่ชัด'),
    ('บรรจุภัณฑ์', 'กล่องบุบระหว่างขนส่ง'),
    ('บรรจุภัณฑ์', 'ซีลฝาขวดไม่แน่น'),
    ('บรรจุภัณฑ์', 'จำนวนในกล่องไม่ครบตามที่ระบุ'),
    ('เอกสารกำกับ', 'ไม่ได้รับใบรับรองผลวิเคราะห์พร้อมสินค้า'),
    ('เอกสารกำกับ', 'เอกสารระบุเลขที่รุ่นผลิตไม่ตรงกับสินค้า'),
    ('การจัดส่ง', 'สินค้าถึงปลายทางช้ากว่ากำหนด'),
    ('การจัดส่ง', 'อุณหภูมิระหว่างขนส่งสูงกว่าเกณฑ์'),
    ('อื่น ๆ', 'สอบถามความคงสภาพหลังเปิดใช้'),
]
REPORTERS = ['โรงพยาบาลรัฐ', 'โรงพยาบาลเอกชน', 'ร้านขายยา', 'ตัวแทนจำหน่ายต่างประเทศ', 'ฝ่ายขายในประเทศ']
SEVERITY = ['ต่ำ', 'ปานกลาง', 'สูง', 'วิกฤต']


def build_complaints():
    rng = random.Random(25680303)
    wb = Workbook()
    ws = wb.active
    ws.title = 'ทะเบียนข้อร้องเรียน'
    ws.append(['เลขที่ข้อร้องเรียน', 'วันที่รับแจ้ง', 'ผลิตภัณฑ์', 'เลขที่รุ่นผลิต',
               'ผู้แจ้ง', 'ประเภทข้อร้องเรียน', 'รายละเอียดโดยย่อ',
               'ระดับความรุนแรง', 'สถานะ', 'วันที่ปิดเรื่อง', 'เลขที่ CAPA'])

    records = []
    seq = 0

    def add(month, day, product, batch, reporter, ctype, detail, severity, status, closed, capa):
        nonlocal seq
        seq += 1
        records.append(['CPL-2568-%03d' % seq, '%02d/%02d/2568' % (day, month), product, batch,
                        reporter, ctype, detail, severity, status, closed, capa])

    # กลุ่มข้อร้องเรียนที่เชื่อมโยงกับยอดขายที่หายไปของ Gentamicin (จุดที่ 3)
    cluster = [
        (6, 11, 'GTM-2568-0416', 'โรงพยาบาลรัฐ', 'สูง'),
        (6, 19, 'GTM-2568-0416', 'โรงพยาบาลเอกชน', 'สูง'),
        (6, 27, 'GTM-2568-0417', 'โรงพยาบาลรัฐ', 'วิกฤต'),
        (7, 4,  'GTM-2568-0417', 'ร้านขายยา', 'วิกฤต'),
        (7, 9,  'GTM-2568-0417', 'ตัวแทนจำหน่ายต่างประเทศ', 'วิกฤต'),
        (7, 16, 'GTM-2568-0418', 'โรงพยาบาลรัฐ', 'วิกฤต'),
        (7, 23, 'GTM-2568-0418', 'โรงพยาบาลเอกชน', 'วิกฤต'),
    ]
    for month, day, batch, reporter, sev in cluster:
        add(month, day, 'Gentamicin T P', batch, reporter, 'คุณภาพทางกายภาพ',
            'พบอนุภาคแขวนลอยในหลอดบรรจุ', sev, 'ปิดเรื่องแล้ว',
            '%02d/08/2568' % min(day + 5, 28), 'CAPA-2568-014')

    # ข้อร้องเรียนทั่วไปที่กระจายทั้งปี
    other_products = [p[0] for p in PRODUCTS if p[0] != 'Gentamicin T P']
    prefix = {name: ''.join(w[0] for w in name.split()[:3]).upper()[:3] for name in other_products}
    while len(records) < 80:
        month = rng.randint(1, 12)
        day = rng.randint(1, 28)
        product = rng.choice(other_products)
        batch = '%s-2568-%04d' % (prefix[product], rng.randint(100, 999))
        ctype, detail = rng.choice(COMPLAINT_TYPES)
        sev = rng.choices(SEVERITY, weights=[42, 38, 16, 4])[0]
        if month <= 10:
            status = rng.choices(['ปิดเรื่องแล้ว', 'อยู่ระหว่างสอบสวน'], weights=[85, 15])[0]
        else:
            status = rng.choices(['ปิดเรื่องแล้ว', 'อยู่ระหว่างสอบสวน'], weights=[45, 55])[0]
        if status == 'ปิดเรื่องแล้ว':
            cm = month + 1 if month < 12 else 12
            closed = '%02d/%02d/2568' % (rng.randint(1, 28), cm)
            capa = 'CAPA-2568-%03d' % rng.randint(1, 40) if sev in ('สูง', 'วิกฤต') else '-'
        else:
            closed = '-'
            capa = 'CAPA-2568-%03d' % rng.randint(1, 40) if sev in ('สูง', 'วิกฤต') else '-'
        add(month, day, product, batch, rng.choice(REPORTERS), ctype, detail, sev, status, closed, capa)

    records.sort(key=lambda r: (int(r[1][3:5]), int(r[1][0:2])))
    for i, rec in enumerate(records, start=1):
        rec[0] = 'CPL-2568-%03d' % i
        ws.append(rec)

    style_header(ws)
    for r in range(2, len(records) + 2):
        for c in range(1, 12):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BOX
            cell.alignment = Alignment(vertical='top', wrap_text=(c == 7))
    autosize(ws, [20, 16, 30, 20, 26, 22, 40, 18, 22, 16, 18])

    add_note_sheet(wb, [
        'ทะเบียนข้อร้องเรียนผลิตภัณฑ์ปี 2568 จำนวน 80 รายการ เรียงตามวันที่รับแจ้ง',
        'ไฟล์นี้ใช้คู่กับ TP-01-sales-monthly.xlsx ในโมดูล M2 — คำตอบของโจทย์ "ยอดตกเพราะอะไร" ต้องอ่านสองไฟล์ประกอบกัน',
        'สถานะ "อยู่ระหว่างสอบสวน" กระจุกตัวในช่วงปลายปีตามธรรมชาติของงาน ไม่ใช่ความผิดปกติ',
    ])
    wb.save(os.path.join(OUT, 'TP-03-complaint-log.xlsx'))
    return len(records)


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    n1 = build_sales()
    build_cost()
    n3 = build_complaints()
    print('TP-01-sales-monthly.xlsx  : %d แถวข้อมูล' % n1)
    print('TP-02-product-cost.xlsx   : %d ผลิตภัณฑ์' % len(PRODUCTS))
    print('TP-03-complaint-log.xlsx  : %d รายการ' % n3)
